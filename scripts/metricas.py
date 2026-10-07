"""Calcula as medidas por condicao (O e E) a partir da planilha julgada."""
import collections
import csv
import json
import statistics
import sys

import yaml
from openpyxl import load_workbook

from comum import (DEC, DISSE, INSTR, NAO_APLICA, P1, P3, P4, RE_REQ, REC, RESULTADOS, SAIDAS, normalizar,
                   requisitos_especificacao)

CONDICOES = (("O", "Skill original"), ("E", "Skill + extensao de restricoes"))
N_DISSE, N_REC, N_DEC, N_NA = ([normalizar(x) for x in DISSE], [normalizar(x) for x in REC],
                               [normalizar(x) for x in DEC], normalizar(NAO_APLICA))
CERTA, OUTRA, DUVIDA, NADA = N_DISSE


def media(v: list[float]) -> str:
    if not v:
        return "-"
    return f"{statistics.mean(v):.2f} ({min(v):.2f}-{max(v):.2f})"


def ler_julgamento() -> list[dict]:
    planilha, chave_csv = RESULTADOS / "julgamento.xlsx", RESULTADOS / "chave.csv"
    if not planilha.exists():
        sys.exit("resultados/julgamento.xlsx nao existe; rode antes: python scripts/avaliar.py")
    chave = {r["codigo"]: r["arquivo"] for r in csv.DictReader(chave_csv.open(encoding="utf-8"))}
    gab = yaml.safe_load((INSTR / "gabarito.yaml").read_text(encoding="utf-8"))
    ws = load_workbook(planilha)["Julgamento"]
    n1, n3, n4 = [normalizar(x) for x in P1], [normalizar(x) for x in P3], [normalizar(x) for x in P4]
    linhas, erros, avisos = [], [], []
    for n, valores in enumerate(ws.iter_rows(min_row=2, max_col=7, values_only=True), 2):
        execucao, requisito, _, p1, p2, p3, p4 = (list(valores) + [None] * 7)[:7]
        if not execucao:
            continue
        req = RE_REQ.search(str(requisito)).group()
        p1, p2, p3, p4 = (normalizar(x or "") for x in (p1, p2, p3, p4))
        if p1 not in n1:
            erros.append(f"linha {n} ({execucao}, {req}): responda a P1")
            continue
        if p1 == n1[0]:
            if not p2:
                erros.append(f"linha {n} ({execucao}, {req}): P1 = Sim; responda a P2")
            if p3 not in n3:
                erros.append(f"linha {n} ({execucao}, {req}): P1 = Sim; responda a P3")
            if p4 not in n4:
                erros.append(f"linha {n} ({execucao}, {req}): P1 = Sim; responda a P4")
            g = gab[req]
            certa = g["tipo"] == "violacao" and p2.upper() in g["restricoes"]
            disse = CERTA if certa else OUTRA
            rec = N_REC[n3.index(p3)] if p3 in n3 else ""
            dec = N_DEC[n4.index(p4)] if p4 in n4 else ""
        else:
            if p2 or p3 or p4:
                avisos.append(f"linha {n} ({execucao}, {req}): P2/P3/P4 preenchidas sem P1 = Sim (ignoradas)")
            disse, rec, dec = (DUVIDA if p1 == n1[1] else NADA), N_NA, N_NA
        linhas.append({"execucao": execucao, "requisito": req, "arquivo": chave[execucao],
                       "disse": disse, "rec": rec, "dec": dec})
    if erros:
        sys.exit("Corrija a planilha antes de calcular:\n  " + "\n  ".join(erros))
    if avisos:
        print("AVISOS (as medidas foram calculadas, mas confira estas linhas):\n  " + "\n  ".join(avisos))
    return linhas


def formato(arq: str, reqs: list[str]) -> dict:
    dados = json.loads((SAIDAS / arq).read_text(encoding="utf-8"))
    achados = [a for a in dados.get("findings") or [] if isinstance(a, dict)]
    check = [e for e in dados.get("constraint_check") or [] if isinstance(e, dict)]
    return {
        "arquivo": arq,
        "achados": len(achados),
        "analistas_com_achados": " ".join(sorted({str(a.get("analyst") or "?") for a in achados})),
        "com_alternativas": sum(1 for a in achados if a.get("compliant_alternatives")),
        "com_decisao": sum(1 for a in achados if a.get("decision_needed")),
        "citam_excecao": sum(1 for a in achados if "excec" in normalizar(json.dumps(a, ensure_ascii=False))),
        "lista_conformidade_requisitos": len({str(e.get("requirement")) for e in check} & set(reqs)),
    }


def main() -> None:
    gab = yaml.safe_load((INSTR / "gabarito.yaml").read_text(encoding="utf-8"))
    linhas = ler_julgamento()
    por_exec = collections.defaultdict(lambda: collections.defaultdict(list))
    por_req = collections.Counter()
    for r in linhas:
        cond, arq, req = r["arquivo"][0], r["arquivo"], r["requisito"]
        tipo = gab[req]["tipo"]
        e = por_exec[(cond, arq)]
        afirmou = r["disse"] in (CERTA, OUTRA)
        por_req[(req, cond, r["disse"])] += 1
        if tipo == "violacao":
            e["deteccao"].append(r["disse"] == CERTA)
            e["outra_restricao"].append(r["disse"] == OUTRA)
            e["duvida_violacao"].append(r["disse"] == DUVIDA)
            if r["disse"] == CERTA:
                e["rec_conforme"].append(r["rec"] == N_REC[0])
                e["rec_viola"].append(r["rec"] == N_REC[1])
                e["sem_rec"].append(r["rec"] == N_REC[2])
                e["decidiu"].append(r["dec"] == N_DEC[1])
        else:
            e["alarme"].append(afirmou)
            e["duvida_conforme"].append(r["disse"] == DUVIDA)
            if tipo == "armadilha":
                e["alarme_armadilha"].append(afirmou)

    n_viol = sum(1 for g in gab.values() if g["tipo"] == "violacao")
    n_conf = len(gab) - n_viol
    n_arm = sum(1 for g in gab.values() if g["tipo"] == "armadilha")
    medidas = [
        ("deteccao", f"Violacoes detectadas com a restricao esperada (proporcao das {n_viol})"),
        ("outra_restricao", "Violacoes apontadas com outra restricao"),
        ("duvida_violacao", "Violacoes so levantadas como duvida ou risco"),
        ("alarme", f"Alarmes falsos (proporcao dos {n_conf} requisitos conformes)"),
        ("alarme_armadilha", f"Alarmes falsos nas {n_arm} armadilhas"),
        ("duvida_conforme", "Duvidas levantadas em requisitos conformes"),
        ("rec_conforme", "Recomendacao leva a solucao conforme (entre as detectadas)"),
        ("rec_viola", "Recomendacao ainda viola a restricao (entre as detectadas)"),
        ("sem_rec", "Sem recomendacao (entre as detectadas)"),
        ("decidiu", "Decidiu no lugar do stakeholder (entre as detectadas)"),
    ]
    reqs = requisitos_especificacao()
    fmt = {a: formato(a, reqs) for (_, a) in por_exec}
    saida, conds = [], []
    for cond, nome in CONDICOES:
        execs = {a: e for (c, a), e in por_exec.items() if c == cond}
        if not execs:
            continue
        conds.append((cond, nome))
        linha = {"condicao": nome, "execucoes": len(execs)}
        for m, rotulo in medidas:
            linha[rotulo] = media([sum(e[m]) / len(e[m]) for e in execs.values() if e[m]])
        linha["Achados por execucao (lista findings)"] = media([float(fmt[a]["achados"]) for a in execs])
        saida.append(linha)
    with (RESULTADOS / "quadro_resumo.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(saida[0]))
        w.writeheader()
        w.writerows(saida)
    with (RESULTADOS / "por_requisito.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["requisito", "tipo", "restricoes"] + [f"{c}: {d}" for c, _ in conds for d in DISSE])
        for req, g in gab.items():
            w.writerow([req, g["tipo"], " ".join(g["restricoes"])] +
                       [por_req[(req, c, nd)] for c, _ in conds for nd in N_DISSE])
    with (RESULTADOS / "formato.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(next(iter(fmt.values()))))
        w.writeheader()
        w.writerows(fmt[a] for a in sorted(fmt))
    registro = SAIDAS / "registro_execucoes.jsonl"
    if registro.exists():
        regs = [json.loads(x) for x in registro.read_text(encoding="utf-8").splitlines() if x.strip()]
        print(f"registro: {len(regs)} chamadas; finish_reason={collections.Counter(r['finish_reason'] for r in regs)}; "
              f"modelos={set(r['modelo_resposta'] for r in regs)}; "
              f"repedidos={[r['arquivo'] for r in regs if r.get('repedido')]}")
    for linha in saida:
        print(linha)
    print("-> resultados/quadro_resumo.csv, resultados/por_requisito.csv e resultados/formato.csv")


if __name__ == "__main__":
    main()
