"""Le as saidas e monta a planilha de julgamento (execucoes codificadas, sem condicao e sem gabarito)."""
import argparse
import csv
import json
import random

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from comum import (P1, P3, P4, PERGUNTAS, RE_REQ, RESULTADOS, SAIDAS, codigos_restricoes, normalizar,
                   requisitos_especificacao, INSTR)

CAMPOS = [
    ("constraint_id", "Restrição"), ("description", "Descrição"), ("issue", "Problema"), ("conflict", "Conflito"),
    ("reasoning", "Justificativa"), ("evidence_requirement", "Trecho do requisito"),
    ("evidence_constraint", "Trecho da restrição"), ("recommendation", "Recomendação"),
    ("compliant_alternatives", "Alternativas"), ("decision_needed", "Decisão a tomar"),
]
CABECALHO = ["execucao", "requisito", "o que a skill escreveu sobre este requisito",
             PERGUNTAS["p1"], PERGUNTAS["p2"], PERGUNTAS["p3"], PERGUNTAS["p4"], "observação (opcional)"]
LARGURAS = [10, 40, 80, 26, 16, 24, 30, 25]


def apontamentos(dados: dict) -> list[dict]:
    vistos, fontes = set(), list(dados.get("findings") or [])
    for grupo in (dados.get("findings_by_principle") or {}).values():
        fontes.extend(grupo or [])
    lista = []
    for a in fontes:
        if not isinstance(a, dict):
            continue
        chave = (str(a.get("analyst", "")), str(a.get("finding_title", "")))
        if chave in vistos:
            continue
        vistos.add(chave)
        lista.append({"dados": a, "reqs": set(RE_REQ.findall(json.dumps(a, ensure_ascii=False))), "lista": False})
    for item in dados.get("constraint_check") or []:
        if isinstance(item, dict) and normalizar(item.get("status", "")) == "violation":
            lista.append({"dados": item, "reqs": set(RE_REQ.findall(str(item.get("requirement", "")))),
                          "lista": True})
    return lista


def texto(a: dict) -> str:
    d = a["dados"]
    if a["lista"]:
        cods = ", ".join(map(str, d.get("constraint_ids") or [])) or "?"
        return f"• Lista final da skill: marcou o requisito como violação de {cods}."
    partes = [f"• {d.get('finding_title') or '(sem título)'} (analista {d.get('analyst') or '?'})"]
    for campo, rotulo in CAMPOS:
        v = d.get(campo)
        if v:
            partes.append(f"   {rotulo}: {'; '.join(map(str, v)) if isinstance(v, list) else v}")
    return "\n".join(partes)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--refazer", action="store_true", help="sobrescreve a planilha de julgamento")
    args = ap.parse_args()
    RESULTADOS.mkdir(exist_ok=True)
    espec = {}
    for linha in (INSTR / "especificacao.md").read_text(encoding="utf-8").splitlines():
        m = RE_REQ.search(linha)
        if m:
            espec[m.group()] = linha.split(":**", 1)[-1].strip()
    reqs = requisitos_especificacao()

    arquivos = sorted(a for a in SAIDAS.glob("[OE]_r*.json") if "tentativa" not in a.name)
    chave_csv = RESULTADOS / "chave.csv"
    if chave_csv.exists():
        chave = {r["arquivo"]: r["codigo"] for r in csv.DictReader(chave_csv.open(encoding="utf-8"))}
    else:
        nomes = [a.name for a in arquivos]
        random.shuffle(nomes)
        chave = {n: f"E{i}" for i, n in enumerate(nomes, 1)}
        with chave_csv.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["arquivo", "codigo"])
            w.writerows(sorted(chave.items()))

    execucoes, todos = [], []
    for arq in arquivos:
        try:
            dados = json.loads(arq.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            print(f"AVISO: {arq.name} nao e JSON valido; fica fora da planilha")
            continue
        lista = apontamentos(dados)
        execucoes.append((chave[arq.name], lista))
        for a in lista:
            todos.append({"arquivo": arq.name, "analista": a["dados"].get("analyst", "lista final"),
                          "titulo": a["dados"].get("finding_title", ""), "requisitos": " ".join(sorted(a["reqs"]))})
    execucoes.sort(key=lambda e: int(e[0][1:]))
    with (RESULTADOS / "achados.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["arquivo", "analista", "titulo", "requisitos"])
        w.writeheader()
        w.writerows(todos)

    planilha = RESULTADOS / "julgamento.xlsx"
    if planilha.exists() and not args.refazer:
        print(f"{planilha.name} ja existe e nao foi sobrescrita (use --refazer)")
        return
    wb = Workbook()
    ajuda = wb.active
    ajuda.title = "Como preencher"
    for linha in COMO_PREENCHER:
        ajuda.append([linha])
    ajuda.column_dimensions["A"].width = 130
    for row in ajuda.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ajuda["A1"].font = Font(bold=True, size=12)

    ws = wb.create_sheet("Julgamento")
    ws.append(CABECALHO)
    amarelo, cinza = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9D9D9")
    for i, c in enumerate(ws[1]):
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = LARGURAS[i]
        if i >= 3:
            c.fill = amarelo
    for codigo, lista in execucoes:
        for req in reqs:
            escritos = [texto(a) for a in lista if req in a["reqs"]]
            ws.append([codigo, f"{req}: {espec.get(req, '')}",
                       ("\n\n".join(escritos) or "(a skill não escreveu nada sobre este requisito)")[:32000]])
    ultima = ws.max_row
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[3:]:
            c.fill = amarelo
    for col, opcoes in (("D", P1), ("E", codigos_restricoes()), ("F", P3), ("G", P4)):
        dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ultima}")
    ws.conditional_formatting.add(f"E2:G{ultima}",
                                  FormulaRule(formula=[f'AND($D2<>"",$D2<>"{P1[0]}")'], fill=cinza))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:H{ultima}"
    wb.active = 1
    wb.save(planilha)
    print(f"{len(execucoes)} execucoes -> {ultima - 1} linhas em resultados/{planilha.name}")


COMO_PREENCHER = [
    "Como preencher a planilha (aba Julgamento)",
    "Cada linha é um requisito em uma execução. As 6 execuções têm códigos sorteados (E1 a E6); não abra "
    "resultados/chave.csv antes de terminar.",
    "A coluna 'o que a skill escreveu' foi copiada da resposta da skill (pasta saidas). Nada foi interpretado e "
    "nenhuma resposta sua vem preenchida.",
    "Leia o que a skill escreveu e responda as perguntas nas colunas amarelas, escolhendo na lista:",
    f"P1. A skill disse que o requisito fere uma restrição? '{P1[0]}': ela afirmou que fere (conflita, viola, contraria). '{P1[1]}': ela só pediu "
    f"para verificar, confirmar, ou falou de risco, sem afirmar que fere. '{P1[2]}': ela não falou de restrições "
    "(pode ter falado de outra coisa, como ambiguidade) ou não escreveu nada.",
    "P2. Qual restrição ela disse que o requisito fere? Escolha o código. Se ela não citar o código, escolha a "
    "restrição de que ela fala.",
    f"P3. A solução que ela sugeriu respeita a restrição? '{P3[0]}': seguir a sugestão resolve o conflito. "
    f"'{P3[1]}': a sugestão continua ferindo a restrição. '{P3[2]}': ela apontou o problema, mas não sugeriu saída.",
    f"P4. Quem decide? '{P4[0]}': ela recomendou uma solução ou deu opções (inclusive pedir exceção), sem reescrever o requisito. '{P4[1]}': ela apresentou o "
    "requisito já reescrito ou tratou a solução como decidida.",
    "Só responda P2, P3 e P4 quando a P1 for 'Sim'. Nas outras linhas, essas células ficam cinza: deixe em branco.",
    "Você não precisa saber se a skill acertou. Depois, o metricas.py compara as suas respostas com o gabarito.",
]


if __name__ == "__main__":
    main()
